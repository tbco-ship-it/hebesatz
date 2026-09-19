#!/usr/bin/env python3
"""Regionaldatenbank Deutschland (GENESIS) scrapes -> data/gemeinden.json
 - hebesatz_B_20241231.json : table 71231-03-01-5, Grundsteuer B Hebesatz, Stand 31.12.2024, all Gemeinden
 - hebesatz_changes_20250630.json : table 71231-02-01-5, Hebesätze changed in H1 2025 (A, B, C, Gewerbe), Stand 30.06.2025
 - hebesaetze-realsteuern.xlsx : Destatis 2022 publication, used for the 2022 rates (fallback for A/Gewerbe) and 2022 population fallback
 - gv_20260331.xlsx : Destatis Gemeindeverzeichnis 31.03.2026 — Kreis names, population 31.12.2024, PLZ, centre coordinates
Current B = H1-2025 value if the change table reports one, else the 31.12.2024 value ('-' in the change table = unchanged on 30.06.2025).
NRW Gemeinden may set separate B rates for residential/non-residential since 2025; the Regionaldatenbank then reports no B — flagged as diff_possible."""
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data/raw"
LAENDER = {"01": ("Schleswig-Holstein", "SH"), "02": ("Hamburg", "HH"), "03": ("Niedersachsen", "NI"), "04": ("Bremen", "HB"), "05": ("Nordrhein-Westfalen", "NW"), "06": ("Hessen", "HE"), "07": ("Rheinland-Pfalz", "RP"), "08": ("Baden-Württemberg", "BW"), "09": ("Bayern", "BY"), "10": ("Saarland", "SL"), "11": ("Berlin", "BE"), "12": ("Brandenburg", "BB"), "13": ("Mecklenburg-Vorpommern", "MV"), "14": ("Sachsen", "SN"), "15": ("Sachsen-Anhalt", "ST"), "16": ("Thüringen", "TH")}
# Grundsteuer model per Land after the 2025 reform (Bundesmodell vs. Ländermodelle)
MODELL = {"BW": "Modifiziertes Bodenwertmodell", "BY": "Flächenmodell", "HH": "Wohnlagenmodell", "HE": "Flächen-Faktor-Modell", "NI": "Flächen-Lage-Modell", "SL": "Bundesmodell mit abweichenden Steuermesszahlen", "SN": "Bundesmodell mit abweichenden Steuermesszahlen"}


def num(v):
    v = (v or "").strip().replace(" ", "").replace(" ", "")
    return int(v) if re.fullmatch(r"\d+", v) else None


def slug(s):
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss").replace("Ä", "Ae").replace("Ö", "Oe").replace("Ü", "Ue")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", s.lower())).strip("-")


def clean_name(n):
    """'Kiel, kreisfreie Stadt, Landeshauptstadt' -> ('Kiel', 'kreisfreie Stadt, Landeshauptstadt')"""
    parts = [x.strip() for x in n.split(",")]
    return parts[0], ", ".join(parts[1:])


def main():
    base = json.loads((RAW / "hebesatz_B_20241231.json").read_text())
    changes = {}
    cp = RAW / "hebesatz_changes_20250630.json"
    if cp.exists():
        for r in json.loads(cp.read_text()):
            changes[r[1]] = {"a": num(r[3]), "b": num(r[4]), "c": num(r[5]), "gew": num(r[6])}
    kreis_name, gv = {}, {}
    ws = openpyxl.load_workbook(RAW / "gv_20260331.xlsx", read_only=True).worksheets[1]
    for row in ws.iter_rows(values_only=True):
        if row[0] == "40" and row[7]:
            kreis_name[f"{row[2]}{row[3]}{row[4]}"] = {"41": "kreisfreie Stadt", "42": "Stadtkreis", "43": "Kreis", "44": "Landkreis", "45": "Regionalverband", "46": "Städteregion"}.get(str(row[1]), "Kreis"), str(row[7]).split(",")[0]
        elif row[0] == "60" and row[7]:
            gv[f"{row[2]}{row[3]}{row[4]}{row[6]}"] = {"ew": int(row[9]) if isinstance(row[9], (int, float)) else None, "plz": str(row[13] or "").strip() or None,
                                                       "lng": float(str(row[14]).replace(",", ".")) if row[14] else None, "lat": float(str(row[15]).replace(",", ".")) if row[15] else None}
    pop, old = {}, {}
    wb = openpyxl.load_workbook(RAW / "hebesaetze-realsteuern.xlsx", read_only=True)
    for ws in wb.worksheets:
        if not ws.title.startswith("Land "):
            continue
        for row in ws.iter_rows(values_only=True):
            ags = str(row[3] or "").strip()
            if re.fullmatch(r"\d{8}", ags):
                p = str(row[5] or "").replace(" ", "")
                pop[ags] = int(p) if p.isdigit() else None
                old[ags] = {"a": row[6] if isinstance(row[6], (int, float)) else None, "b": row[7] if isinstance(row[7], (int, float)) else None, "gew": row[8] if isinstance(row[8], (int, float)) else None}
    gem = []
    seen = defaultdict(int)
    for code, name, val, note in base:
        b24 = num(val)
        ch = changes.get(code) or {}
        b25 = ch.get("b")
        if b24 is None and b25 is None:
            continue  # gemeindefreies Gebiet / no rate
        short, suffix = clean_name(name)
        land = LAENDER[code[:2]]
        s = slug(short)
        seen[(code[:2], s)] += 1
        if seen[(code[:2], s)] > 1:
            s = f"{s}-{seen[(code[:2], s)]}"
        o = old.get(code, {})
        k = kreis_name.get(code[:5], ("Kreis", ""))
        v = gv.get(code, {})
        gem.append({"ags": code, "name": short, "suffix": suffix, "land": land[0], "land_code": land[1], "land_slug": slug(land[0]), "slug": s, "kreis": code[:5],
                    "kreis_name": k[1], "kreis_type": k[0], "plz": v.get("plz"), "lat": v.get("lat"), "lng": v.get("lng"),
                    "b2024": b24, "b2025": b25, "a2025": ch.get("a"), "c2025": ch.get("c"), "gew2025": ch.get("gew"),
                    "b2022": int(o["b"]) if o.get("b") is not None else None, "a2022": int(o["a"]) if o.get("a") is not None else None, "gew2022": int(o["gew"]) if o.get("gew") is not None else None,
                    "ew": v.get("ew") if v.get("ew") is not None else pop.get(code), "ew_stand": "31.12.2024" if v.get("ew") is not None else "30.06.2022",
                    "reported_2025": b25 is not None, "changed_2025": b25 is not None and b24 is not None and b25 != b24,
                    # NRW: A re-set for 2025 but no B reported → the Gemeinde most likely set separate Wohn-/Nichtwohn rates the table cannot hold
                    "diff_possible": land[1] == "NW" and b25 is None and ch.get("a") is not None})
    for g in gem:
        g["b"] = g["b2025"] if g["b2025"] is not None else g["b2024"]
        g["stand"] = "30.06.2025"
    out = {"source": {"name": "Regionaldatenbank Deutschland (Statistische Ämter des Bundes und der Länder), Tabellen 71231-03-01-5 und 71231-02-01-5; Destatis Hebesätze der Realsteuern 2022; Destatis Gemeindeverzeichnis 31.03.2026 (Bevölkerung 31.12.2024)",
                      "url": "https://www.regionalstatistik.de/genesis/online?operation=table&code=71231-03-01-5", "fetched": "2026-09-19", "license": "Datenlizenz Deutschland – Namensnennung – Version 2.0"},
           "modell": MODELL, "gemeinden": gem}
    (ROOT / "data/gemeinden.json").write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")))
    print(f"{len(gem)} Gemeinden, {sum(1 for g in gem if g['reported_2025'])} reported for H1 2025 ({sum(1 for g in gem if g['changed_2025'])} changed), {sum(1 for g in gem if g['diff_possible'])} NRW diff_possible, {sum(1 for g in gem if g['ew'])} with population, {sum(1 for g in gem if g['kreis_name'])} with Kreis")


if __name__ == "__main__":
    main()
